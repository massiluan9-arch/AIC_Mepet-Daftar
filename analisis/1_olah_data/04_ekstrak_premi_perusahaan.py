"""Sebaran premi harta benda penutupan langsung per perusahaan asuransi umum, 2024.

Sumber: OJK, Statistik Perasuransian 2024, Tabel 2.11 (premi per perusahaan per lini usaha).
Dipakai untuk memberi skala rupiah bagi PT Aktu Proteksi (perusahaan fiktif): hasil model dinyatakan
per 100 premi, lalu dikalikan premi harta benda perusahaan ukuran menengah.
Keluaran: data/premi_harta_benda_per_perusahaan.json

Jalankan: python -I 04_ekstrak_premi_perusahaan.py
"""
import json
import statistics
import sys
import warnings
from pathlib import Path

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")
DI_SINI = Path(__file__).resolve().parent
BERKAS = DI_SINI.parents[1] / "data" / "industri_asuransi" / "ojk_statistik_2024" / "2.11 Premi Usaha Asuransi Umum dan Reasuransi Tahun 2024.xlsx"
KELUAR = DI_SINI / "data" / "premi_harta_benda_per_perusahaan.json"


def main():
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(BERKAS, read_only=True, data_only=True)
    lembar = None
    for ws in wb.worksheets:
        kepala = " ".join(str(c) for r in ws.iter_rows(min_row=1, max_row=9, values_only=True) for c in r if isinstance(c, str))
        if "Harta Benda" in kepala:
            lembar = ws
            break
    assert lembar is not None, "lembar lini harta benda tidak ditemukan"
    premi, bagian = {}, None
    for r in lembar.iter_rows(values_only=True):
        teks = r[1] if isinstance(r[1], str) else ""
        if "REASURANSI" in teks.upper() and "PERUSAHAAN" in teks.upper():
            bagian = "reasuransi"
        elif "PERUSAHAAN ASURANSI UMUM" in teks.upper():
            bagian = "umum"
        if bagian == "umum" and isinstance(r[0], (int, float)) and isinstance(r[2], (int, float)) and teks.startswith("PT"):
            premi[teks.strip()] = float(r[2])
    aktif = sorted(v for v in premi.values() if v > 0)
    kuartil = statistics.quantiles(aktif, n=4)
    out = {
        "sumber": "OJK, Statistik Perasuransian 2024, Tabel 2.11, lini harta benda, kolom pos langsung",
        "lembar": lembar.title.strip(), "satuan": "juta rupiah",
        "jumlah_perusahaan": len(premi), "jumlah_menulis_harta_benda": len(aktif),
        "jumlah_premi": sum(aktif), "rata": statistics.mean(aktif), "median": statistics.median(aktif),
        "kuartil_1": kuartil[0], "kuartil_3": kuartil[2], "terbesar": aktif[-1],
        "porsi_lima_terbesar": sum(aktif[-5:]) / sum(aktif),
    }
    KELUAR.write_text(json.dumps(out, indent=1, ensure_ascii=False), encoding="utf-8")
    for k, v in out.items():
        print(f"{k}: {v:,.1f}" if isinstance(v, float) else f"{k}: {v}")


if __name__ == "__main__":
    main()
