"""Besaran hasil underwriting lini harta benda dari statistik OJK.

Sumber:
  1. OJK, Statistik Perasuransian 2024, Tabel 2.10.4 (lini harta benda): premi, komisi, klaim pos langsung.
  2. OJK, Statistik Perasuransian 2024, Key Indicator: rasio klaim harta benda 2011-2024 (neto) dan beban
     manajemen industri.
  3. OJK, Statistik Asuransi Juli 2025, laba rugi asuransi umum konvensional: beban usaha.
Keluaran: data/underwriting_industri.json

Jalankan: python -I 03_ekstrak_underwriting.py
"""
import json
import statistics
import sys
import warnings
from pathlib import Path

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")
DI_SINI = Path(__file__).resolve().parent
AKAR = DI_SINI.parents[1]
DIR_OJK = AKAR / "data" / "industri_asuransi" / "ojk_statistik_2024"
BERKAS_2025 = AKAR / "data" / "industri_asuransi" / "ojk_statistik_2025" / "OJK_Statistik_Asuransi_Juli_2025.xlsx"
KELUAR = DI_SINI / "data" / "underwriting_industri.json"


def buka(berkas):
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return openpyxl.load_workbook(berkas, read_only=True, data_only=True)


def tabel_2_10_4():
    """Kolom pos langsung lini harta benda, juta rupiah."""
    ws = buka(DIR_OJK / "2.10 Rekapitulasi Premi, Klaim dan Komisi Tahun 2024.xlsx")["4 REKAP"]
    out = {}
    for r in ws.iter_rows(values_only=True):
        if isinstance(r[1], str) and isinstance(r[2], (int, float)):
            out[r[1].split("/")[0].strip()] = float(r[2])
    return out


def key_indicator():
    """Deret 2011-2024: rasio klaim harta benda (%), premi bruto dan beban manajemen (triliun rupiah)."""
    ws = buka(DIR_OJK / "0. Key Indicator 2024.xlsx")["CETAK"]
    baris = list(ws.iter_rows(values_only=True))
    tahun = list(range(2011, 2025))

    def deret(i):
        nilai = [c for c in baris[i] if isinstance(c, (int, float))]
        assert len(nilai) == len(tahun), (i, nilai)
        return dict(zip(tahun, nilai))

    i_rasio = next(i for i, r in enumerate(baris) if any(isinstance(c, str) and "Incurred Loss Ratios" in c for c in r))
    i_harta = next(i for i in range(i_rasio, i_rasio + 5) if any(isinstance(c, str) and c.strip() == "Property" for c in baris[i]))
    i_premi = next(i for i, r in enumerate(baris) if any(isinstance(c, str) and c.startswith("Gross Premium (Trillion") for c in r))
    i_beban = next(i for i, r in enumerate(baris) if any(isinstance(c, str) and c.strip() == "Management Expense" for c in r))
    return deret(i_harta), deret(i_premi), deret(i_beban)


def laba_rugi_2025():
    ws = buka(BERKAS_2025)["IS-General Insurance"]
    out = {}
    for r in ws.iter_rows(values_only=True):
        if isinstance(r[0], (int, float)) and isinstance(r[1], str) and isinstance(r[8], (int, float)):
            out[r[1].strip()] = float(r[8])
    return out


def main():
    t = tabel_2_10_4()
    premi, komisi = t["Premi Dicatat"], t["Komisi"]
    premi_diterima, klaim_terjadi = t["Premi Diterima"], t["Klaim yang Terjadi"]
    rasio, premi_bruto, beban = key_indicator()
    lr = laba_rugi_2025()
    premi_2025 = lr["Premi Penutupan Langsung"]

    nilai = list(rasio.values())
    rasio_beban = {th: beban[th] / premi_bruto[th] for th in rasio}
    # Simpangan rasio beban di sekitar garis tren (selisih tahunan dibagi akar dua)
    selisih = [rasio_beban[th] - rasio_beban[th - 1] for th in list(rasio)[1:]]
    out = {
        "sumber": "OJK, Statistik Perasuransian 2024 (Tabel 2.10.4 dan Key Indicator); OJK, Statistik Asuransi Juli 2025",
        "rasio_klaim_harta_benda_neto_persen": rasio,
        "rasio_klaim_rata": statistics.mean(nilai) / 100,
        "rasio_klaim_simpangan": statistics.stdev(nilai) / 100,
        "rasio_klaim_terendah": min(nilai) / 100,
        "rasio_klaim_tertinggi": max(nilai) / 100,
        "rasio_klaim_langsung_bruto_2024": klaim_terjadi / premi_diterima,
        "komisi_langsung_2024": komisi / premi,
        "beban_usaha_per_premi_langsung_2025": lr["Jumlah Beban Usaha"] / premi_2025,
        "komisi_semua_lini_2025": -lr["Komisi Dibayar"] / premi_2025,
        "klaim_bruto_semua_lini_per_premi_langsung_2025": lr["Klaim Bruto"] / premi_2025,
        "rasio_beban_manajemen_per_premi_bruto": rasio_beban,
        "simpangan_rasio_beban": statistics.stdev(selisih) / 2 ** 0.5,
        "catatan": {
            "rasio_klaim": "Deret 2011-2024 adalah rasio klaim neto (setelah reasuransi) dan satu-satunya deret panjang yang "
                           "diterbitkan. Rasio pos langsung bruto hanya tersedia untuk 2024.",
            "beban": "Beban usaha = pemasaran, pegawai dan pengurus, pendidikan, umum dan administrasi; tujuh bulan 2025; "
                     "pembilang dan penyebut sama-sama tujuh bulan.",
        },
    }
    KELUAR.write_text(json.dumps(out, indent=1, ensure_ascii=False), encoding="utf-8")
    for k in ("rasio_klaim_rata", "rasio_klaim_simpangan", "rasio_klaim_langsung_bruto_2024", "komisi_langsung_2024",
              "beban_usaha_per_premi_langsung_2025", "komisi_semua_lini_2025", "simpangan_rasio_beban"):
        print(f"{k}: {out[k]:.4f}")


if __name__ == "__main__":
    main()
