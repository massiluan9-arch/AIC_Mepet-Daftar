"""Menghitung ulang buku kerja lewat Excel, lalu memeriksa galat rumus dan mencocokkan hasil kunci.

openpyxl menulis rumus tanpa nilai tersimpan. Skrip ini membuka berkas di Excel (COM), menghitung
ulang, menyimpan, lalu membaca nilai tersimpan untuk (1) memastikan tidak ada sel bergalat dan
(2) mencocokkan beberapa hasil dengan hitungan Python dari rancangan.json.

Pemakaian: python 5_solusi_motera/kode/hitung_ulang_excel.py
"""
import json
import sys
from pathlib import Path

import openpyxl
import win32com.client
import yaml

sys.stdout.reconfigure(encoding="utf-8")
DIR_G = Path(__file__).resolve().parents[1]
BERKAS = DIR_G / "Analisis_Solusi_AIC.xlsx"
GALAT = ("#NAME?", "#DIV/0!", "#VALUE!", "#REF!", "#N/A", "#NUM!", "#NULL!")


def hitung_ulang():
    excel = win32com.client.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    try:
        wb = excel.Workbooks.Open(str(BERKAS))
        excel.CalculateFullRebuild()
        wb.Save()
        wb.Close(False)
    finally:
        excel.Quit()


def main():
    hitung_ulang()
    nilai = openpyxl.load_workbook(BERKAS, data_only=True)
    rumus = openpyxl.load_workbook(BERKAS)
    jumlah_rumus, galat, kosong = 0, [], []
    for ws in rumus.worksheets:
        for baris in ws.iter_rows():
            for c in baris:
                if isinstance(c.value, str) and c.value.startswith("="):
                    jumlah_rumus += 1
                    v = nilai[ws.title][c.coordinate].value
                    if v is None:
                        kosong.append(f"{ws.title}!{c.coordinate}")
                    elif isinstance(v, str) and v in GALAT:
                        galat.append(f"{ws.title}!{c.coordinate} {v}")
    print(f"rumus: {jumlah_rumus}; galat: {len(galat)}; tanpa nilai: {len(kosong)}")
    for g in galat[:20] + kosong[:20]:
        print("  ", g)

    # pencocokan dengan hitungan Python
    h = json.loads((DIR_G / "keluaran" / "rancangan.json").read_text(encoding="utf-8"))
    ps = yaml.safe_load((DIR_G / "parameter_solusi.yaml").read_text(encoding="utf-8"))
    p, k, n = h["perbandingan"], ps["biaya_modal"]["nilai"], ps["horizon_tahun"]
    skala = ps["skala_rupiah"]["premi_harta_benda_miliar"] / 100
    dE = (p["modal_saja"]["ekuitas"] - p["terpilih"]["ekuitas"]) * skala
    dL = (p["modal_saja"]["laba_bersih_rata"] - p["terpilih"]["laba_bersih_rata"]) * skala
    dD = (p["modal_saja"]["defisit_harapan"] - p["terpilih"]["defisit_harapan"]) * skala
    awal = sum(v["nilai"] for v in ps["biaya_implementasi"]["awal"].values())
    tahunan = sum(v["nilai"] for v in ps["biaya_implementasi"]["tahunan"].values())
    arus = [dE - awal] + [-(dL - dD + tahunan)] * (n - 1) + [-(dL - dD + tahunan) - dE]
    npv = sum(a / (1 + k) ** t for t, a in enumerate(arus))
    bm = nilai["Biaya_Manfaat"]
    # Rev. 2: perspektif perusahaan dan biaya pelaksanaan dua kali, dihitung ulang di Python untuk pencocokan
    arus_p = [dE - awal] + [-(dL + tahunan)] * (n - 1) + [-(dL + tahunan) - dE]
    npv_p = sum(a / (1 + k) ** t for t, a in enumerate(arus_p))
    arus_x2 = [dE - 2 * awal] + [-(dL - dD + 2 * tahunan)] * (n - 1) + [-(dL - dD + 2 * tahunan) - dE]
    npv_x2 = sum(a / (1 + k) ** t for t, a in enumerate(arus_x2))
    arus_p_x2 = [dE - 2 * awal] + [-(dL + 2 * tahunan)] * (n - 1) + [-(dL + 2 * tahunan) - dE]
    npv_p_x2 = sum(a / (1 + k) ** t for t, a in enumerate(arus_p_x2))
    kd = h.get("kerugian_dialihkan", {})
    periksa = [("NPV", bm["B14"].value, npv), ("NPV rumus tertutup", nilai["Sensitivitas"]["B15"].value, npv),
               ("Premi dilepas", nilai["Matematika_Keuangan"]["C6"].value, p["terpilih"]["premi_dilepas"]),
               ("Laba ekonomi solusi", nilai["Hasil_Model"]["D24"].value, p["terpilih"]["laba_ekonomi"]),
               ("Nilai bersih solusi", nilai["Hasil_Model"]["D25"].value, p["terpilih"]["nilai_bersih"]),
               ("NPV perspektif perusahaan", bm["B51"].value, npv_p), ("NPV diperluas, biaya dua kali", bm["B55"].value, npv_x2),
               ("NPV perusahaan, biaya dua kali", bm["B57"].value, npv_p_x2),
               ("Kerugian dialihkan rata-rata (Rp miliar)", bm["B60"].value, kd.get("rata_per_tahun", 0.0) * skala)]
    lulus = not galat and not kosong
    for nama, excel_, python_ in periksa:
        cocok = excel_ is not None and abs(excel_ - python_) < 1e-6 * max(1, abs(python_))
        lulus &= cocok
        print(f"{'LULUS' if cocok else 'GAGAL'}  {nama}: Excel {excel_} lawan Python {python_}")
    hasil = {"npv_miliar": bm["B14"].value, "biaya_setara_modal": bm["B15"].value, "nilai_kini_manfaat": bm["B17"].value,
             "nilai_kini_biaya": bm["B18"].value, "rasio_manfaat_biaya": bm["B19"].value, "keputusan": bm["B20"].value,
             "biaya_keamanan_miliar": bm["B33"].value, "biaya_keamanan_terhadap_premi": bm["B34"].value,
             "biaya_per_0_1_poin": bm["B35"].value, "ekuitas_dihemat_miliar": dE, "laba_dilepas_miliar": dL,
             "defisit_dihindari_miliar": dD, "biaya_awal_miliar": awal, "biaya_tahunan_miliar": tahunan,
             "selisih_rumus_simulasi": nilai["Matematika_Keuangan"]["C14"].value,
             "porsi_laba_ditahan": nilai["Matematika_Keuangan"]["C18"].value,
             "npv_per_biaya_modal": {str(nilai["Sensitivitas"][f"A{i}"].value): nilai["Sensitivitas"][f"B{i}"].value for i in range(12, 19)},
             "npv_per_horizon": {str(nilai["Sensitivitas"][f"A{i}"].value): nilai["Sensitivitas"][f"B{i}"].value for i in range(22, 26)},
             "npv_per_kasus": {nilai["Sensitivitas"][f"A{i}"].value: nilai["Sensitivitas"][f"J{i}"].value for i in range(29, 34)},
             "kelayakan": {nilai["Kelayakan"][f"A{i}"].value: {"skor": nilai["Kelayakan"][f"C{i}"].value, "putusan": nilai["Kelayakan"][f"D{i}"].value}
                           for i in range(21, 26) if nilai["Kelayakan"][f"A{i}"].value},
             "jumlah_rumus": jumlah_rumus, "galat_rumus": len(galat),
             # ---- Rev. 2 (kunci baru; kunci lama tidak diubah) ----
             "profil_acuan": h["meta"]["profil"], "aturan_peluang": h["meta"].get("aturan_peluang"),
             "strategi_terpilih": h["terpilih"]["strategi"], "target_terpilih": h["terpilih"]["target"],
             "target_modal_saja": p["modal_saja"]["target"],
             "npv_diperluas_miliar": bm["B14"].value, "npv_perusahaan_miliar": bm["B51"].value,
             "npv_diperluas_biaya_x2_miliar": bm["B55"].value, "npv_perusahaan_biaya_x2_miliar": bm["B57"].value,
             "impas_biaya_modal_diperluas": bm["B53"].value, "impas_biaya_modal_perusahaan": bm["B52"].value,
             "impas_biaya_modal_diperluas_biaya_x2": bm["B58"].value, "impas_biaya_modal_perusahaan_biaya_x2": bm["B59"].value,
             "kerugian_dialihkan_rata_miliar": bm["B60"].value, "kerugian_dialihkan_ekor_0_5_miliar": bm["B61"].value,
             "porsi_kerugian_ekor_dialihkan": bm["B62"].value, "premi_tidak_dibayar_tertanggung_miliar": bm["B63"].value,
             "skala_premi_syarat_terpilih": nilai["Asumsi"]["B12"].value, "porsi_premi_terkena_syarat": nilai["Asumsi"]["B11"].value}
    (DIR_G / "keluaran" / "hasil_excel.json").write_text(json.dumps(hasil, indent=1, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(hasil, indent=1, ensure_ascii=False)[:1800])
    sys.exit(0 if lulus else 1)


if __name__ == "__main__":
    main()
