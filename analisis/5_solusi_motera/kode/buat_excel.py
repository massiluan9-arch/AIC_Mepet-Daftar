"""Menyusun buku kerja Excel analisis solusi: matematika keuangan, biaya-manfaat, kepekaan, dan kerangka bisnis.

Angka model dibaca dari keluaran/rancangan.json dan 4_simulasi_terpadu/keluaran/hasil_terpadu.json.
Semua turunan di dalam buku kerja ditulis sebagai rumus Excel supaya ikut berubah bila asumsi diganti.
Konvensi warna: biru = masukan atau keluaran simulasi; hitam = rumus; hijau = rujukan ke lembar lain.

Pemakaian: python 5_solusi_motera/kode/buat_excel.py
Keluaran: 5_solusi_motera/Analisis_Solusi_AIC.xlsx
"""
import json
import sys
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference, ScatterChart, Series
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding="utf-8")
DIR_G = Path(__file__).resolve().parents[1]
DIR_T = DIR_G.parent
BERKAS = DIR_G / "Analisis_Solusi_AIC.xlsx"

HURUF = "Arial"
BIRU, HITAM, HIJAU = "0000FF", "000000", "008000"
F_BIASA = Font(name=HURUF, size=10, color=HITAM)
F_MASUK = Font(name=HURUF, size=10, color=BIRU)
F_TAUT = Font(name=HURUF, size=10, color=HIJAU)
F_TEBAL = Font(name=HURUF, size=10, bold=True)
F_JUDUL = Font(name=HURUF, size=13, bold=True)
F_KECIL = Font(name=HURUF, size=9, italic=True, color="555555")
ISI_KEPALA = PatternFill("solid", fgColor="D9E1F2")
ISI_KUNCI = PatternFill("solid", fgColor="FFFF00")
ISI_HASIL = PatternFill("solid", fgColor="E2EFDA")
TIPIS = Side(style="thin", color="999999")
KOTAK = Border(left=TIPIS, right=TIPIS, top=TIPIS, bottom=TIPIS)
BUNGKUS = Alignment(wrap_text=True, vertical="top")
ANGKA = '#,##0.00;(#,##0.00);"-"'
ANGKA1 = '#,##0.0;(#,##0.0);"-"'
PERSEN = '0.00%;(0.00%);"-"'
PERSEN1 = '0.0%'


def tulis(ws, sel, nilai, font=F_BIASA, format_angka=None, isi=None, bungkus=False, kotak=False, rata=None):
    c = ws[sel]
    c.value = nilai
    c.font = font
    if format_angka:
        c.number_format = format_angka
    if isi:
        c.fill = isi
    if bungkus:
        c.alignment = BUNGKUS
    if rata:
        c.alignment = Alignment(horizontal=rata, vertical="top", wrap_text=bungkus)
    if kotak:
        c.border = KOTAK
    return c


def kepala(ws, baris, judul, kolom_awal=1):
    for i, t in enumerate(judul):
        c = ws.cell(row=baris, column=kolom_awal + i, value=t)
        c.font, c.fill, c.border = F_TEBAL, ISI_KEPALA, KOTAK
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")


def lebar(ws, ukuran):
    for i, w in enumerate(ukuran, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def idn(x, d=1):
    """Angka dengan koma desimal untuk teks berbahasa Indonesia."""
    return f"{x:.{d}f}".replace(".", ",")


def judul(ws, teks, keterangan=None):
    tulis(ws, "A1", teks, F_JUDUL)
    if keterangan:
        tulis(ws, "A2", keterangan, F_KECIL)


def main():
    h = json.loads((DIR_G / "keluaran" / "rancangan.json").read_text(encoding="utf-8"))
    t = json.loads((DIR_T / "4_simulasi_terpadu" / "keluaran" / "hasil_terpadu.json").read_text(encoding="utf-8"))
    ps = yaml.safe_load((DIR_G / "parameter_solusi.yaml").read_text(encoding="utf-8"))
    pt = yaml.safe_load((DIR_T / "4_simulasi_terpadu" / "parameter_terpadu.yaml").read_text(encoding="utf-8"))
    meta_b = json.loads((DIR_T / "2_model_bencana" / "keluaran" / "04_meta.json").read_text(encoding="utf-8"))
    st_pilih, s_pilih = h["terpilih"]["strategi"], h["terpilih"]["target"]
    p = h["perbandingan"]
    calon = h["calon"][st_pilih]
    # Rev. 2: skala premi dibaca dari strategi terpilih, bukan ditulis mati untuk batas 10%.
    # porsi = porsi premi gempa yang terkena syarat khusus (1 bila syarat berlaku untuk seluruh portofolio);
    # skala efektif = 1 - (1 - faktor premi gempa strategi) / porsi, sehingga premi dilepas = premi gempa x porsi x (1 - skala).
    khusus = _strategi(st_pilih, "syarat_khusus")
    khusus = khusus if isinstance(khusus, str) and khusus else ""
    porsi_syarat = float(_strategi(st_pilih, "porsi_premi_khusus")) if khusus else 1.0
    skala_syarat = 1.0 - (1.0 - calon["faktor_premi_gempa"]) / porsi_syarat
    nama_syarat = khusus or str(_strategi(st_pilih, "syarat"))
    kd = h.get("kerugian_dialihkan", {})
    wb = Workbook()

    # ------------------------------------------------------------------ Baca_Dulu
    ws = wb.active
    ws.title = "Baca_Dulu"
    judul(ws, "Analisis solusi PT Aktu Proteksi: Geo-Solvency Shield")
    baris = [
        ("Isi", "Buku kerja ini memuat perhitungan matematika keuangan, analisis biaya-manfaat, kepekaan, skenario, dan kerangka bisnis untuk solusi soal 5."),
        ("Satuan", "Angka model dinyatakan per 100 premi dicatat lini harta benda setahun. Lembar Asumsi mengubahnya ke rupiah lewat skala premi."),
        ("Warna", "Biru = masukan atau keluaran simulasi yang boleh diganti; hitam = rumus; hijau = rujukan ke lembar lain; kuning = asumsi kunci."),
        ("Asal angka model", "Lembar Hasil_Model disalin dari 5_solusi_motera/keluaran/rancangan.json (simulasi 100.000 tahun; R untuk kerugian bencana, Python untuk neraca)."),
        ("Solusi terpilih", f"{calon['keterangan']}; rasio solvabilitas awal {idn(100 * s_pilih, 0)}%"),
        ("Tiga jalan", "Sekarang = polis standar pada target 120%. Modal saja = polis standar pada target yang lolos tiga kendala. Solusi = rancangan terpilih."),
        ("Tiga kendala", "C1: P(R < 100%) paling besar 0,5% (kurva campuran). C2: sama, pada kurva Aulady & Fujimi saja. C3: rasio paling rendah 100% setelah skenario megathrust Jawa bagian barat M 8,7 (guncangan median)."),
        ("Kerangka bisnis", "Dipilih menurut aturan satu pertanyaan satu kerangka: 5 Whys (akar masalah), Impact-Effort (prioritas), NPV dan kepekaan (kelayakan finansial), matriks kelayakan empat aspek, Risk Matrix (risiko), Roadmap dan KPI (pelaksanaan)."),
        ("Lembar", "Asumsi, Hasil_Model, Matematika_Keuangan, Biaya_Manfaat, Sensitivitas, Skenario, Lima_Mengapa, Dampak_Upaya, Kelayakan, Risiko, Peta_Jalan_KPI."),
    ]
    for i, (a, b) in enumerate(baris, 3):
        tulis(ws, f"A{i}", a, F_TEBAL, bungkus=True)
        tulis(ws, f"B{i}", b, bungkus=True)
    lebar(ws, [20, 120])

    # ------------------------------------------------------------------ Asumsi
    ws = wb.create_sheet("Asumsi")
    judul(ws, "Asumsi dan masukan", "Sel biru boleh diganti; lembar lain mengikuti.")
    kepala(ws, 3, ["Butir", "Nilai", "Satuan", "Status", "Sumber atau catatan"])
    bm = h["biaya_manfaat"]
    impl = ps["biaya_implementasi"]
    isi = [
        ("Premi harta benda setahun", ps["skala_rupiah"]["premi_harta_benda_miliar"], "Rp miliar", "asumsi", ps["skala_rupiah"]["catatan"], ANGKA1, True),
        ("Skala: Rp miliar per satuan per 100 premi", "=B4/100", "Rp miliar", "rumus", "Pengali dari satuan model ke rupiah", ANGKA, False),
        ("Biaya modal", ps["biaya_modal"]["nilai"], "per tahun", "asumsi", ps["biaya_modal"]["catatan"], PERSEN1, True),
        ("Horizon analisis", ps["horizon_tahun"], "tahun", "asumsi", "Lembar Biaya_Manfaat disusun untuk lima tahun", "0", False),
        ("Tarif pajak penghasilan badan", pt["underwriting"]["pajak"]["nilai"], "", "regulasi", "UU 7/2021", PERSEN1, False),
        ("Eksposur gempa per premi harta benda (E)", pt["bencana"]["eksposur_per_premi"]["nilai"], "kali", "data", "MAIPARK (2025); OJK (2025)", ANGKA, False),
        ("Tarif gempa rata-rata", pt["bencana"]["tarif_gempa_per_mil"]["nilai"], "per mil", "data", "MAIPARK (2025)", '0.0000', False),
        ("Porsi premi gempa yang terkena syarat strategi terpilih", porsi_syarat, "", "data", f"{calon['keterangan']}; zona inti: Papua, Jawa Barat, Jakarta, Sumatera Utara, Banten; MAIPARK (2025)", PERSEN1, False),
        (f"Skala premi syarat strategi terpilih ({nama_syarat})", skala_syarat, "", "regulasi", "SEOJK 6/2017 Lampiran III Tabel III.B; dibaca dari strategi terpilih", PERSEN1, False),
        ("Komisi", pt["underwriting"]["komisi"]["nilai"], "dari premi", "data", "OJK Tabel 2.10.4", PERSEN1, False),
        ("Defisit harapan pemegang polis ikut dihitung (1 = ya, 0 = tidak)", 1, "", "asumsi", "Butsic (1994); Myers & Read (2001)", "0", True),
    ]
    for i, (a, b, c, d, e, f, kunci) in enumerate(isi, 4):
        tulis(ws, f"A{i}", a, bungkus=True)
        tulis(ws, f"B{i}", b, F_BIASA if isinstance(b, str) else F_MASUK, f, ISI_KUNCI if kunci else None)
        tulis(ws, f"C{i}", c)
        tulis(ws, f"D{i}", d)
        tulis(ws, f"E{i}", e, bungkus=True)
    r = 4 + len(isi) + 1                    # baris 16
    tulis(ws, f"A{r}", "Biaya pelaksanaan awal (Rp miliar)", F_TEBAL)
    awal = list(impl["awal"].items())
    for i, (k, v) in enumerate(awal, r + 1):
        tulis(ws, f"A{i}", k.replace("_", " ").capitalize())
        tulis(ws, f"B{i}", v["nilai"], F_MASUK, ANGKA1)
        tulis(ws, f"C{i}", "Rp miliar"); tulis(ws, f"D{i}", "asumsi"); tulis(ws, f"E{i}", v["catatan"], bungkus=True)
    r_awal = r + 1 + len(awal)              # baris 20
    tulis(ws, f"A{r_awal}", "Jumlah biaya awal", F_TEBAL)
    tulis(ws, f"B{r_awal}", f"=SUM(B{r + 1}:B{r_awal - 1})", F_TEBAL, ANGKA1)
    tulis(ws, f"A{r_awal + 1}", "Biaya operasi tahunan (Rp miliar)", F_TEBAL)
    tahunan = list(impl["tahunan"].items())
    for i, (k, v) in enumerate(tahunan, r_awal + 2):
        tulis(ws, f"A{i}", k.replace("_", " ").capitalize())
        tulis(ws, f"B{i}", v["nilai"], F_MASUK, ANGKA1)
        tulis(ws, f"C{i}", "Rp miliar"); tulis(ws, f"D{i}", "asumsi"); tulis(ws, f"E{i}", v["catatan"], bungkus=True)
    r_thn = r_awal + 2 + len(tahunan)       # baris 25
    tulis(ws, f"A{r_thn}", "Jumlah biaya tahunan", F_TEBAL)
    tulis(ws, f"B{r_thn}", f"=SUM(B{r_awal + 2}:B{r_thn - 1})", F_TEBAL, ANGKA1)
    tulis(ws, f"A{r_thn + 1}", "Pengali biaya pelaksanaan (uji)", F_BIASA)
    tulis(ws, f"B{r_thn + 1}", 1, F_MASUK, "0.0", ISI_KUNCI)
    tulis(ws, f"E{r_thn + 1}", "Isi 2 untuk menguji biaya dua kali lipat", bungkus=True)
    lebar(ws, [52, 14, 12, 11, 78])
    A = {"premi": "Asumsi!$B$4", "skala": "Asumsi!$B$5", "k": "Asumsi!$B$6", "n": "Asumsi!$B$7", "pajak": "Asumsi!$B$8",
         "E": "Asumsi!$B$9", "tarif": "Asumsi!$B$10", "porsi": "Asumsi!$B$11", "skala10": "Asumsi!$B$12", "komisi": "Asumsi!$B$13",
         "pakai_defisit": "Asumsi!$B$14", "awal": f"Asumsi!$B${r_awal}", "tahunan": f"Asumsi!$B${r_thn}", "kali": f"Asumsi!$B${r_thn + 1}"}

    # ------------------------------------------------------------------ Hasil_Model
    ws = wb.create_sheet("Hasil_Model")
    judul(ws, "Keluaran simulasi untuk tiga jalan", "Per 100 premi dicatat, kecuali disebut lain. Sel biru = keluaran simulasi 100.000 tahun.")
    kepala(ws, 3, ["Ukuran", "Sekarang", "Modal saja", "Solusi", "Keterangan"])
    urut = ["sekarang", "modal_saja", "terpilih"]
    ukuran = [
        ("Strategi", "strategi", None, "Kode strategi pada model bencana"),
        ("Rasio solvabilitas awal", "target", PERSEN1, "Tingkat Solvabilitas dibagi MMBR"),
        ("Ekuitas", "ekuitas", ANGKA, "Dana pemegang saham yang tertanam"),
        ("MMBR", "mmbr", ANGKA, "SEOJK 24/2017, lima modul"),
        ("Premi", "premi", ANGKA, "Setelah premi gempa yang dilepas"),
        ("Hasil underwriting non-bencana, rata-rata", "U_rata", ANGKA, ""),
        ("Hasil investasi, rata-rata", "I_rata", ANGKA, ""),
        ("Kerugian bencana, rata-rata", "L_rata", ANGKA, ""),
        ("Kerugian bencana, VaR 99,5%", "L_var_99_5", ANGKA, "Kerugian setahun yang terlampaui 1 dari 200 tahun"),
        ("Kerugian bencana, TVaR 99,5%", "L_tvar_99_5", ANGKA, "Rata-rata kerugian pada 0,5% tahun terburuk"),
        ("Klaim satu kejadian, VaR 99,5%", "klaim_kejadian_var_99_5", ANGKA, "Dasar ukuran kantong likuiditas"),
        ("Laba bersih harapan", "laba_bersih_rata", ANGKA, "Setelah pajak"),
        ("P(R < 100%), kurva campuran", "peluang_di_bawah_100", PERSEN, "Kendala C1: paling besar 0,50%"),
        ("P(R < 100%), kurva Aulady & Fujimi", "peluang_aulady", PERSEN, "Kendala C2: paling besar 0,50%"),
        ("P(R < 0%)", "peluang_di_bawah_0", PERSEN, "AYD lebih kecil daripada liabilitas"),
        ("Defisit harapan pemegang polis", "defisit_harapan", ANGKA, "Rata-rata kekurangan AYD terhadap liabilitas"),
        ("Rasio setelah skenario megathrust", "rasio_setelah_megathrust", PERSEN1, "Kendala C3: paling rendah 100%"),
        ("Aset siap cair", "aset_siap_cair", ANGKA, "Kas, deposito, SBN"),
    ]
    R = {}
    for i, (nama, kunci, fmt, ket) in enumerate(ukuran, 4):
        R[kunci] = i
        tulis(ws, f"A{i}", nama)
        for j, jalan in enumerate(urut):
            tulis(ws, f"{'BCD'[j]}{i}", p[jalan][kunci], F_MASUK, fmt)
        tulis(ws, f"E{i}", ket, bungkus=True)
    r0 = 4 + len(ukuran) + 1
    tulis(ws, f"A{r0}", "Turunan (rumus)", F_TEBAL)
    turunan = [
        ("Laba ekonomi = laba bersih - biaya modal x ekuitas", "laba_ekonomi", "={c}%d-%s*{c}%d" % (R["laba_bersih_rata"], A["k"], R["ekuitas"]), ANGKA),
        ("Nilai bersih = laba ekonomi - defisit harapan", "nilai_bersih", "={c}%d-%s*{c}%d" % (r0 + 1, A["pakai_defisit"], R["defisit_harapan"]), ANGKA),
        ("Periode ulang pelanggaran (tahun)", "periode", "=1/{c}%d" % R["peluang_di_bawah_100"], "0"),
        ("Rasio klaim gempa harapan terhadap premi gempa penuh", "rasio_gempa", "={c}%d/(%s*%s/10)" % (R["L_rata"], A["E"], A["tarif"]), PERSEN1),
        ("Ekuitas (Rp miliar)", "ekuitas_rp", "={c}%d*%s" % (R["ekuitas"], A["skala"]), ANGKA1),
        ("Laba bersih harapan (Rp miliar)", "laba_rp", "={c}%d*%s" % (R["laba_bersih_rata"], A["skala"]), ANGKA1),
        ("Kerugian 1 dalam 200 tahun (Rp miliar)", "var_rp", "={c}%d*%s" % (R["L_var_99_5"], A["skala"]), ANGKA1),
    ]
    T = {}
    for i, (nama, kunci, rumus, fmt) in enumerate(turunan, r0 + 1):
        T[kunci] = i
        tulis(ws, f"A{i}", nama)
        for c in "BCD":
            tulis(ws, f"{c}{i}", rumus.format(c=c), F_BIASA, fmt, ISI_HASIL if kunci in ("laba_ekonomi", "nilai_bersih") else None)
    lebar(ws, [52, 15, 15, 15, 60])

    # ------------------------------------------------------------------ Matematika_Keuangan
    ws = wb.create_sheet("Matematika_Keuangan")
    judul(ws, "Matematika keuangan: dari syarat polis ke laba dan likuiditas", "Per 100 premi dicatat. Rumus menelusuri angka model langkah demi langkah.")
    kepala(ws, 3, ["No", "Langkah", "Nilai", "Rumus dalam kata"])
    HM = "Hasil_Model!"
    langkah = [
        ("Premi gempa pada polis standar", f"={A['E']}*{A['tarif']}/10", "E x tarif per mil / 1.000 x 100", ANGKA, F_TAUT),
        ("Premi gempa yang terkena syarat strategi terpilih", f"=C4*{A['porsi']}", "(1) x porsi premi yang terkena syarat", ANGKA, F_TAUT),
        (f"Premi yang dilepas karena syarat ({calon['keterangan']})", f"=C5*(1-{A['skala10']})", "(2) x (1 - skala premi loss limit)", ANGKA, F_TAUT),
        ("Komisi yang ikut turun", f"=C6*{A['komisi']}", "(3) x komisi", ANGKA, F_TAUT),
        ("Kerugian bencana harapan yang turun", f"={HM}B{R['L_rata']}-{HM}D{R['L_rata']}", "Rata-rata kerugian sekarang - solusi", ANGKA, F_TAUT),
        ("Hasil underwriting sebelum pajak yang dilepas", "=C6-C7-C8", "(3) - (4) - (5)", ANGKA, F_BIASA),
        ("Setelah pajak", f"=C9*(1-{A['pajak']})", "(6) x (1 - tarif pajak)", ANGKA, F_TAUT),
        ("Tambahan hasil investasi atas tambahan ekuitas, setelah pajak", f"=({HM}D{R['I_rata']}-{HM}B{R['I_rata']})*(1-{A['pajak']})", "Selisih hasil investasi x (1 - tarif pajak)", ANGKA, F_TAUT),
        ("Laba bersih yang dilepas menurut rumus", "=C10-C11", "(7) - (8)", ANGKA, F_BIASA),
        ("Laba bersih yang dilepas menurut simulasi", f"={HM}B{R['laba_bersih_rata']}-{HM}D{R['laba_bersih_rata']}", "Laba bersih sekarang - solusi", ANGKA, F_TAUT),
        ("Selisih rumus dan simulasi", "=C12-C13", "Timbul karena pajak hanya dikenakan pada tahun laba", ANGKA, F_BIASA),
    ]
    for i, (nama, rumus, kata, fmt, font) in enumerate(langkah, 4):
        tulis(ws, f"A{i}", i - 3); tulis(ws, f"B{i}", nama, bungkus=True)
        tulis(ws, f"C{i}", rumus, font, fmt); tulis(ws, f"D{i}", kata, bungkus=True)
    tulis(ws, "A16", "Penyangga modal", F_TEBAL)
    penyangga = [
        ("Tambahan ekuitas dari 120% ke target solusi", f"={HM}D{R['ekuitas']}-{HM}B{R['ekuitas']}", "Ekuitas solusi - ekuitas sekarang", ANGKA),
        ("Porsi laba bersih setahun yang perlu ditahan", f"=C17/{HM}D{R['laba_bersih_rata']}", "(12) / laba bersih solusi", PERSEN1),
        ("Ekuitas yang dihemat dibanding jalan modal saja", f"={HM}C{R['ekuitas']}-{HM}D{R['ekuitas']}", "Ekuitas modal saja - ekuitas solusi", ANGKA),
        ("Ekuitas solusi sebagai porsi ekuitas modal saja", f"={HM}D{R['ekuitas']}/{HM}C{R['ekuitas']}", "", PERSEN1),
        ("Tambahan ekuitas (Rp miliar)", f"=C17*{A['skala']}", "x skala", ANGKA1),
        ("Ekuitas yang dihemat (Rp miliar)", f"=C19*{A['skala']}", "x skala", ANGKA1),
    ]
    for i, (nama, rumus, kata, fmt) in enumerate(penyangga, 17):
        tulis(ws, f"A{i}", i - 5); tulis(ws, f"B{i}", nama, bungkus=True)
        tulis(ws, f"C{i}", rumus, F_TAUT, fmt); tulis(ws, f"D{i}", kata, bungkus=True)
    tulis(ws, "A24", "Kantong likuiditas bencana: kas dan deposito lawan klaim satu kejadian 1 dalam 200 tahun", F_TEBAL)
    kepala(ws, 25, ["", "Varian profil Konservatif", "Kas", "Deposito", "Kas + deposito", "Klaim satu kejadian", "Cakupan", "Cukup?"])
    for i, (kode, v) in enumerate(h["terpilih"]["kantong_likuiditas"].items(), 26):
        tulis(ws, f"B{i}", {"K80": "SBN 80%, deposito 20%", "K75": "SBN 75%, deposito 25%", "K70": "SBN 70%, deposito 30%"}[kode])
        tulis(ws, f"C{i}", v["kas"], F_MASUK, ANGKA); tulis(ws, f"D{i}", v["deposito"], F_MASUK, ANGKA)
        tulis(ws, f"E{i}", f"=C{i}+D{i}", F_BIASA, ANGKA)
        tulis(ws, f"F{i}", f"={HM}D{R['klaim_kejadian_var_99_5']}", F_TAUT, ANGKA)
        tulis(ws, f"G{i}", f"=E{i}/F{i}", F_BIASA, PERSEN1)
        tulis(ws, f"H{i}", f'=IF(G{i}>=1,"Ya","Belum")')
    tulis(ws, "B30", "Kas dan deposito dicairkan tanpa menjual SBN. Varian yang cakupannya paling sedikit 100% menjadi pilihan alokasi.", F_KECIL)
    lebar(ws, [6, 58, 15, 44, 15, 18, 11, 10])

    # ------------------------------------------------------------------ Biaya_Manfaat
    ws = wb.create_sheet("Biaya_Manfaat")
    judul(ws, "Analisis biaya-manfaat", "Rp miliar pada skala premi di lembar Asumsi.")
    tulis(ws, "A3", "A. Solusi lawan jalan modal saja (dua cara mencapai tingkat keamanan yang sama)", F_TEBAL)
    kepala(ws, 4, ["Arus kas inkremental", "Tahun 0", "Tahun 1", "Tahun 2", "Tahun 3", "Tahun 4", "Tahun 5", "Keterangan"])
    dE = f"({HM}C{R['ekuitas']}-{HM}D{R['ekuitas']})*{A['skala']}"
    dL = f"({HM}C{R['laba_bersih_rata']}-{HM}D{R['laba_bersih_rata']})*{A['skala']}"
    dD = f"({HM}C{R['defisit_harapan']}-{HM}D{R['defisit_harapan']})*{A['skala']}*{A['pakai_defisit']}"
    baris_bm = [
        ("Ekuitas yang tidak perlu disetor", [f"={dE}", 0, 0, 0, 0, f"=-{dE}"], "Disetor di awal pada jalan modal saja, kembali di akhir horizon"),
        ("Laba bersih yang dilepas", [0] + [f"=-{dL}"] * 5, "Premi gempa dilepas dan hasil investasi atas modal yang tidak ditahan"),
        ("Defisit harapan pemegang polis yang dihindari", [0] + [f"={dD}"] * 5, "Dihitung bila sakelar di lembar Asumsi bernilai 1"),
        ("Biaya pelaksanaan awal", [f"=-{A['awal']}*{A['kali']}", 0, 0, 0, 0, 0], "Mesin akumulasi, penyesuaian produk, pelatihan"),
        ("Biaya operasi tahunan", [0] + [f"=-{A['tahunan']}*{A['kali']}"] * 5, "Tim, data, validasi"),
    ]
    for i, (nama, nilai, ket) in enumerate(baris_bm, 5):
        tulis(ws, f"A{i}", nama, bungkus=True)
        for j, v in enumerate(nilai):
            tulis(ws, f"{'BCDEFG'[j]}{i}", v, F_TAUT if isinstance(v, str) else F_BIASA, ANGKA1)
        tulis(ws, f"H{i}", ket, bungkus=True)
    tulis(ws, "A10", "Arus kas bersih", F_TEBAL)
    tulis(ws, "A11", "Faktor diskonto")
    tulis(ws, "A12", "Nilai kini")
    for j, c in enumerate("BCDEFG"):
        tulis(ws, f"{c}10", f"=SUM({c}5:{c}9)", F_TEBAL, ANGKA1)
        tulis(ws, f"{c}11", f"=1/(1+{A['k']})^{j}", F_TAUT, "0.0000")
        tulis(ws, f"{c}12", f"={c}10*{c}11", F_BIASA, ANGKA1)
    ringkas = [
        ("Nilai kini bersih (NPV)", "=SUM(B12:G12)", ANGKA1, "Positif berarti solusi lebih murah daripada modal saja"),
        ("Biaya setara modal dari solusi", "=IRR(B10:G10,0.1)", PERSEN, "Tingkat diskonto yang membuat NPV nol; bandingkan dengan biaya modal"),
        ("Biaya modal (asumsi)", f"={A['k']}", PERSEN, ""),
        ("Nilai kini manfaat", f"=B5+NPV({A['k']},C7:G7)", ANGKA1, "Ekuitas tidak disetor dan defisit yang dihindari"),
        ("Nilai kini biaya", f"=-(B8+NPV({A['k']},C6:G6)+NPV({A['k']},C9:G9)+G5*G11)", ANGKA1, "Laba dilepas, biaya pelaksanaan, pengembalian ekuitas"),
        ("Rasio manfaat terhadap biaya", "=B17/B18", "0.000", "Di atas 1 berarti manfaat lebih besar"),
        ("Keputusan", '=IF(B14>0,"Solusi lebih murah daripada modal saja","Modal saja lebih murah; lihat lembar Sensitivitas")', None, ""),
    ]
    for i, (nama, rumus, fmt, ket) in enumerate(ringkas, 14):
        tulis(ws, f"A{i}", nama, F_TEBAL)
        tulis(ws, f"B{i}", rumus, F_TEBAL, fmt, ISI_HASIL)
        tulis(ws, f"H{i}", ket, bungkus=True)

    tulis(ws, "A23", "B. Solusi lawan keadaan sekarang (yang tidak memenuhi ambang)", F_TEBAL)
    kepala(ws, 24, ["Butir", "Sekarang", "Solusi", "Perubahan", "", "", "", "Keterangan"])
    banding = [
        ("Laba ekonomi setahun (Rp miliar)", f"={HM}B{T['laba_ekonomi']}*{A['skala']}", f"={HM}D{T['laba_ekonomi']}*{A['skala']}", ANGKA1, "Laba bersih dikurangi biaya modal"),
        ("Ekuitas (Rp miliar)", f"={HM}B{T['ekuitas_rp']}", f"={HM}D{T['ekuitas_rp']}", ANGKA1, "Tambahan dapat dibentuk dari laba ditahan"),
        ("P(R < 100%), kurva campuran", f"={HM}B{R['peluang_di_bawah_100']}", f"={HM}D{R['peluang_di_bawah_100']}", PERSEN, "Ambang 0,50%"),
        ("P(R < 100%), kurva Aulady & Fujimi", f"={HM}B{R['peluang_aulady']}", f"={HM}D{R['peluang_aulady']}", PERSEN, "Ambang 0,50%"),
        ("Kerugian bencana TVaR 99,5% (Rp miliar)", f"={HM}B{R['L_tvar_99_5']}*{A['skala']}", f"={HM}D{R['L_tvar_99_5']}*{A['skala']}", ANGKA1, ""),
        ("Defisit harapan pemegang polis (Rp miliar)", f"={HM}B{R['defisit_harapan']}*{A['skala']}", f"={HM}D{R['defisit_harapan']}*{A['skala']}", ANGKA1, ""),
        ("Rasio setelah skenario megathrust", f"={HM}B{R['rasio_setelah_megathrust']}", f"={HM}D{R['rasio_setelah_megathrust']}", PERSEN1, "Paling rendah 100%"),
    ]
    for i, (nama, a, b, fmt, ket) in enumerate(banding, 25):
        tulis(ws, f"A{i}", nama, bungkus=True)
        tulis(ws, f"B{i}", a, F_TAUT, fmt); tulis(ws, f"C{i}", b, F_TAUT, fmt)
        tulis(ws, f"D{i}", f"=C{i}-B{i}", F_BIASA, fmt)
        tulis(ws, f"H{i}", ket, bungkus=True)
    tulis(ws, "A33", "Biaya keamanan setahun (Rp miliar)", F_TEBAL)
    tulis(ws, "B33", f"=-D25+{A['tahunan']}*{A['kali']}+PMT({A['k']},{A['n']},-{A['awal']}*{A['kali']})", F_TEBAL, ANGKA1, ISI_HASIL)
    tulis(ws, "H33", "Laba ekonomi yang dilepas, biaya operasi, dan anuitas biaya awal", bungkus=True)
    tulis(ws, "A34", "Biaya keamanan terhadap premi", F_TEBAL)
    tulis(ws, "B34", f"=B33/{A['premi']}", F_TEBAL, PERSEN, ISI_HASIL)
    tulis(ws, "A35", "Biaya per 0,1 poin penurunan P(R < 100%) (Rp miliar)", F_TEBAL)
    tulis(ws, "B35", "=B33/(-D27*1000)", F_TEBAL, ANGKA1, ISI_HASIL)

    tulis(ws, "A38", "C. Pembentukan penyangga dari laba ditahan", F_TEBAL)
    kepala(ws, 39, ["Butir (Rp miliar)", "Tahun 1", "Tahun 2", "Tahun 3", "Tahun 4", "Tahun 5", "", "Keterangan"])
    tulis(ws, "A40", "Ekuitas awal tahun")
    tulis(ws, "A41", "Laba bersih harapan")
    tulis(ws, "A42", "Laba ditahan untuk penyangga")
    tulis(ws, "A43", "Laba yang dapat dibagikan")
    tulis(ws, "A44", "Ekuitas akhir tahun")
    tulis(ws, "A45", "Ekuitas sasaran")
    for j, c in enumerate("BCDEF"):
        tulis(ws, f"{c}40", f"={HM}B{T['ekuitas_rp']}" if j == 0 else f"={'BCDEF'[j - 1]}44", F_TAUT if j == 0 else F_BIASA, ANGKA1)
        tulis(ws, f"{c}41", f"={HM}D{T['laba_rp']}", F_TAUT, ANGKA1)
        tulis(ws, f"{c}45", f"={HM}D{T['ekuitas_rp']}", F_TAUT, ANGKA1)
        tulis(ws, f"{c}42", f"=MAX(0,MIN({c}41,{c}45-{c}40))", F_BIASA, ANGKA1)
        tulis(ws, f"{c}43", f"={c}41-{c}42", F_BIASA, ANGKA1)
        tulis(ws, f"{c}44", f"={c}40+{c}42", F_BIASA, ANGKA1)
    tulis(ws, "H42", "Penyangga dibentuk dulu, baru laba dibagikan", bungkus=True)
    # ---- Rev. 2: perspektif perusahaan, biaya pelaksanaan dua kali, titik impas, dan kerugian yang dialihkan
    tulis(ws, "A48", "D. Perspektif perusahaan, uji biaya dua kali, dan kerugian yang dialihkan (Rev. 2)", F_TEBAL)
    kepala(ws, 49, ["Butir", "Tahun 0", "Tahun 1", "Tahun 2", "Tahun 3", "Tahun 4", "Tahun 5", "Keterangan"])
    tulis(ws, "A50", "Arus kas bersih, perspektif perusahaan", bungkus=True)
    tulis(ws, "H50", "Arus kas bersih tanpa suku defisit harapan pemegang polis (bukan arus kas pemegang saham)", bungkus=True)
    tulis(ws, "A54", "Arus kas bersih diperluas, biaya pelaksanaan dua kali", bungkus=True)
    tulis(ws, "A56", "Arus kas bersih perusahaan, biaya pelaksanaan dua kali", bungkus=True)
    for c in "BCDEFG":
        tulis(ws, f"{c}50", f"={c}10-{c}7", F_BIASA, ANGKA1)
        tulis(ws, f"{c}54", f"={c}10+({c}8+{c}9)*(2/{A['kali']}-1)", F_BIASA, ANGKA1)
        tulis(ws, f"{c}56", f"={c}54-{c}7", F_BIASA, ANGKA1)
    tambahan = [
        (51, "NPV perspektif perusahaan (tanpa suku defisit)", "=SUMPRODUCT(B50:G50,B11:G11)", ANGKA1, "Positif berarti solusi lebih murah daripada modal saja bagi pemegang saham"),
        (52, "Titik impas biaya modal, perspektif perusahaan", "=IRR(B50:G50,0.1)", PERSEN, "Di atas titik ini solusi lebih murah"),
        (53, "Titik impas biaya modal, diperluas", "=B15", PERSEN, "Sama dengan baris 15"),
        (55, "NPV diperluas, biaya pelaksanaan dua kali", "=SUMPRODUCT(B54:G54,B11:G11)", ANGKA1, "Uji pengali biaya 2 (parameter_solusi.yaml)"),
        (57, "NPV perspektif perusahaan, biaya pelaksanaan dua kali", "=SUMPRODUCT(B56:G56,B11:G11)", ANGKA1, ""),
        (58, "Titik impas biaya modal diperluas, biaya dua kali", "=IRR(B54:G54,0.1)", PERSEN, ""),
        (59, "Titik impas biaya modal perusahaan, biaya dua kali", "=IRR(B56:G56,0.1)", PERSEN, ""),
        (60, "Kerugian dialihkan ke pemegang polis, rata-rata setahun (Rp miliar)", f"={kd.get('rata_per_tahun', 0)}*{A['skala']}", ANGKA1,
         "Kerugian polis standar dikurangi kerugian strategi terpilih pada tahun simulasi yang sama"),
        (61, "Kerugian dialihkan pada 0,5% tahun terburuk (Rp miliar)", f"={kd.get('rata_ekor_0_5', 0)}*{A['skala']}", ANGKA1, "Rata-rata pada tahun ekor polis standar"),
        (62, "Porsi kerugian ekor polis standar yang dialihkan", kd.get("porsi_ekor_0_5", 0), PERSEN1, ""),
        (63, "Premi yang tidak lagi dibayar tertanggung setahun (Rp miliar)", f"={kd.get('premi_dilepas', 0)}*{A['skala']}", ANGKA1, "Premi gempa yang dilepas"),
    ]
    for i, nama, rumus, fmt, ket in tambahan:
        tulis(ws, f"A{i}", nama, F_TEBAL, bungkus=True)
        tulis(ws, f"B{i}", rumus, F_TEBAL if isinstance(rumus, str) else F_MASUK, fmt, ISI_HASIL)
        tulis(ws, f"H{i}", ket, bungkus=True)
    lebar(ws, [50, 14, 14, 14, 14, 14, 14, 62])
    diagram = BarChart()
    diagram.type = "col"
    diagram.title = "Arus kas inkremental solusi lawan modal saja (Rp miliar)"
    diagram.add_data(Reference(ws, min_col=1, max_col=7, min_row=10, max_row=10), from_rows=True, titles_from_data=True)
    diagram.set_categories(Reference(ws, min_col=2, max_col=7, min_row=4, max_row=4))
    diagram.legend = None
    diagram.height, diagram.width = 7.5, 16
    ws.add_chart(diagram, "J3")

    # ------------------------------------------------------------------ Sensitivitas
    ws = wb.create_sheet("Sensitivitas")
    judul(ws, "Kepekaan nilai kini bersih (solusi lawan modal saja)", "Rp miliar. Rumus tertutup: NPV = dE x (1 - v^n) - (dLaba - dDefisit + biaya tahunan) x a(n,k) - biaya awal.")
    tulis(ws, "A3", "Masukan dasar (dari lembar lain)", F_TEBAL)
    dasar = [("Ekuitas dihemat (dE)", f"={dE}"), ("Laba bersih dilepas setahun (dLaba)", f"={dL}"),
             ("Defisit dihindari setahun (dDefisit)", f"={dD}"), ("Biaya operasi tahunan", f"={A['tahunan']}*{A['kali']}"),
             ("Biaya awal", f"={A['awal']}*{A['kali']}")]
    for i, (nama, rumus) in enumerate(dasar, 4):
        tulis(ws, f"A{i}", nama); tulis(ws, f"B{i}", rumus, F_TAUT, ANGKA1)
    rumus_npv = "={dE}*(1-(1+{k})^-{n})-({dL}-{dD}+{th})*(1-(1+{k})^-{n})/{k}-{aw}"
    tulis(ws, "A10", "1. Biaya modal", F_TEBAL)
    kepala(ws, 11, ["Biaya modal", "NPV", "Keputusan"])
    for i, k in enumerate([0.10, 0.11, 0.12, 0.13, 0.14, 0.15, 0.16], 12):
        tulis(ws, f"A{i}", k, F_MASUK, PERSEN1)
        tulis(ws, f"B{i}", rumus_npv.format(dE="$B$4", dL="$B$5", dD="$B$6", th="$B$7", aw="$B$8", k=f"A{i}", n=A["n"]), F_BIASA, ANGKA1)
        tulis(ws, f"C{i}", f'=IF(B{i}>0,"Solusi lebih murah","Modal saja lebih murah")')
    tulis(ws, "A20", "2. Horizon", F_TEBAL)
    kepala(ws, 21, ["Horizon (tahun)", "NPV", "Keputusan"])
    for i, n in enumerate([3, 5, 7, 10], 22):
        tulis(ws, f"A{i}", n, F_MASUK, "0")
        tulis(ws, f"B{i}", rumus_npv.format(dE="$B$4", dL="$B$5", dD="$B$6", th="$B$7", aw="$B$8", k=A["k"], n=f"A{i}"), F_BIASA, ANGKA1)
        tulis(ws, f"C{i}", f'=IF(B{i}>0,"Solusi lebih murah","Modal saja lebih murah")')
    tulis(ws, "A27", "3. Asumsi model (target tiap jalan dihitung ulang oleh simulasi)", F_TEBAL)
    kepala(ws, 28, ["Kasus", "Target modal saja", "Target solusi", "Ekuitas modal saja", "Ekuitas solusi", "Laba modal saja", "Laba solusi",
                    "Defisit modal saja", "Defisit solusi", "NPV (Rp miliar)", "Ekuitas solusi / modal saja"])
    for i, b in enumerate(bm["kasus_ulang"], 29):
        tulis(ws, f"A{i}", b["kasus"])
        for c, kunci, fmt in (("B", "target_modal_saja", PERSEN1), ("C", "target_terpilih", PERSEN1), ("D", "ekuitas_modal_saja", ANGKA),
                              ("E", "ekuitas_terpilih", ANGKA), ("F", "laba_modal_saja", ANGKA), ("G", "laba_terpilih", ANGKA),
                              ("H", "defisit_modal_saja", ANGKA), ("I", "defisit_terpilih", ANGKA)):
            tulis(ws, f"{c}{i}", b.get(kunci), F_MASUK, fmt)
        tulis(ws, f"J{i}", rumus_npv.format(dE=f"((D{i}-E{i})*{A['skala']})", dL=f"((F{i}-G{i})*{A['skala']})",
                                            dD=f"((H{i}-I{i})*{A['skala']}*{A['pakai_defisit']})", th="$B$7", aw="$B$8", k=A["k"], n=A["n"]),
              F_BIASA, ANGKA1, ISI_HASIL)
        tulis(ws, f"K{i}", f"=E{i}/D{i}", F_BIASA, PERSEN1)
    r_akhir = 29 + len(bm["kasus_ulang"])
    tulis(ws, f"A{r_akhir + 1}", "Titik impas biaya modal (NPV = 0) pada kasus dasar", F_TEBAL)
    tulis(ws, f"B{r_akhir + 1}", "=Biaya_Manfaat!B15", F_TAUT, PERSEN, ISI_HASIL)
    tulis(ws, f"A{r_akhir + 2}", "Di atas titik ini solusi lebih murah daripada menambah modal; di bawahnya modal lebih murah.", F_KECIL)
    lebar(ws, [44, 16, 26, 16, 16, 16, 16, 16, 16, 16, 18])

    # ------------------------------------------------------------------ Skenario
    ws = wb.create_sheet("Skenario")
    judul(ws, "Uji skenario gempa (Scenario Analysis)", "Kerugian satu kejadian pada kurva campuran; underwriting dan investasi pada nilai harapan.")
    kepala(ws, 3, ["Skenario", "Guncangan", "Jalan", "Rasio awal", "Kerugian (per 100 premi)", "Kerugian (Rp miliar)", "Rasio akhir", "Kekurangan kas (per 100 premi)", "Status"])
    for i, b in enumerate(h["terpilih"]["skenario"], 4):
        tulis(ws, f"A{i}", b["skenario"]); tulis(ws, f"B{i}", "median" if b["eta"] == 0 else "median + 1 simpangan baku")
        tulis(ws, f"C{i}", "Sekarang" if b["strategi"] == "S0_dasar" else "Solusi")
        tulis(ws, f"D{i}", b["target"], F_MASUK, PERSEN1)
        tulis(ws, f"E{i}", b["kerugian"], F_MASUK, ANGKA1)
        tulis(ws, f"F{i}", f"=E{i}*{A['skala']}", F_TAUT, ANGKA1)
        tulis(ws, f"G{i}", b["rasio_akhir"], F_MASUK, PERSEN1)
        tulis(ws, f"H{i}", b["kekurangan_kas"], F_MASUK, ANGKA1)
        tulis(ws, f"I{i}", f'=IF(G{i}>=1.2,"Di atas target minimum",IF(G{i}>=1,"Memenuhi 100%",IF(G{i}>=0,"Di bawah 100%","AYD di bawah liabilitas")))')
    lebar(ws, [38, 26, 12, 12, 16, 16, 12, 18, 26])

    # ------------------------------------------------------------------ Lima_Mengapa
    ws = wb.create_sheet("Lima_Mengapa")
    judul(ws, "Akar masalah: 5 Whys", "Dipilih karena ada satu rantai sebab yang dominan. Angka dari Hasil_Model dan simulasi.")
    s0 = p["sekarang"]
    m_asuransi = s0["modul"]["asuransi"]
    mengapa = [
        ("Masalah", f"Pada target 120% peluang rasio solvabilitas jatuh di bawah 100% dalam setahun {idn(100 * s0['peluang_di_bawah_100'], 2)}%, di atas ambang 0,5%."),
        ("Mengapa 1", f"Kerugian bencana 1 dalam 200 tahun {idn(s0['L_var_99_5'], 0)} per 100 premi, sedangkan penyangga 20% di atas MMBR hanya {idn(0.2 * s0['mmbr'])} dan laba bersih setahun {idn(s0['laba_bersih_rata'], 0)}."),
        ("Mengapa 2", "Seluruh eksposur gempa ditahan sendiri dan polis standar mengganti sampai nilai pertanggungan penuh; yang dipotong hanya risiko sendiri 2,5%."),
        ("Mengapa 3", f"Eksposur terkumpul. Lima provinsi memuat {idn(100 * meta_b['porsi_eksposur_zona_inti'], 0)}% eksposur dan sekitar 85% kerugian tahun ekor; {idn(100 * meta_b['porsi_jumlah_risiko_besar'])}% risiko memuat {idn(100 * meta_b['porsi_eksposur_risiko_besar'], 0)}% eksposur."),
        ("Mengapa 4", f"MMBR menghitung risiko asuransi dari cadangan (25% CAPYBMP, 20% cadangan klaim), bukan dari eksposur bencana. Modul asuransi {idn(m_asuransi)} per 100 premi."),
        ("Mengapa 5", "Belum ada batas kerugian kontraktual di zona yang menumpuk, dan besar eksposur yang boleh ditahan belum dikaitkan dengan modal."),
        ("Akar masalah", "Eksposur gempa yang terkumpul dan tidak berbatas belum terhubung dengan ukuran modal."),
        ("Lalu apa", "Solusi harus membatasi kerugian per risiko di zona inti dan menetapkan target modal dari uji peluang serta uji skenario, bukan dari angka minimum regulasi."),
    ]
    for i, (a, b) in enumerate(mengapa, 4):
        tulis(ws, f"A{i}", a, F_TEBAL); tulis(ws, f"B{i}", b, bungkus=True)
    lebar(ws, [16, 130])

    # ------------------------------------------------------------------ Dampak_Upaya
    ws = wb.create_sheet("Dampak_Upaya")
    judul(ws, "Prioritas strategi mitigasi: Impact-Effort Matrix", "Dampak = penurunan P(R < 100%) pada target 120%; upaya = laba bersih yang dilepas. Keduanya keluaran simulasi.")
    kepala(ws, 3, ["Strategi", "Keterangan", "Upaya: laba dilepas (per 100 premi)", "P(R < 100%)", "Dampak: penurunan peluang (poin persen)", "Kuadran"])
    s2 = t["soal2"]
    p0 = s2["S0_dasar"]["peluang_di_bawah_100"]
    daftar = [st for st in s2 if st != "S0_dasar"]
    akhir = 3 + len(daftar)
    for i, st in enumerate(daftar, 4):
        tulis(ws, f"A{i}", st); tulis(ws, f"B{i}", s2[st]["keterangan"], bungkus=True)
        tulis(ws, f"C{i}", s2[st]["laba_dilepas"], F_MASUK, ANGKA)
        tulis(ws, f"D{i}", s2[st]["peluang_di_bawah_100"], F_MASUK, PERSEN)
        tulis(ws, f"E{i}", f"=({p0}-D{i})*100", F_BIASA, '0.000')
        tulis(ws, f"F{i}", f'=IF(E{i}>=$C${akhir + 3},IF(C{i}<=$C${akhir + 2},"Kerjakan dulu","Proyek besar"),IF(C{i}<=$C${akhir + 2},"Tambahan cepat","Hindari"))')
    tulis(ws, f"A{akhir + 2}", "Batas upaya (median)", F_TEBAL); tulis(ws, f"C{akhir + 2}", f"=MEDIAN(C4:C{akhir})", F_BIASA, ANGKA)
    tulis(ws, f"A{akhir + 3}", "Batas dampak (median)", F_TEBAL); tulis(ws, f"C{akhir + 3}", f"=MEDIAN(E4:E{akhir})", F_BIASA, '0.000')
    tulis(ws, f"A{akhir + 5}", "Batasan: risiko sendiri di atas 2,5% tampak tanpa biaya karena tabel tarif tidak memberi potongan premi; strategi itu tidak dijadikan calon.", F_KECIL)
    lebar(ws, [32, 52, 20, 14, 22, 18])
    sebar = ScatterChart()
    sebar.title = "Dampak lawan upaya"
    sebar.style = 13
    sebar.x_axis.title = "Upaya: laba bersih dilepas (per 100 premi)"
    sebar.y_axis.title = "Dampak: penurunan peluang (poin persen)"
    seri = Series(Reference(ws, min_col=5, min_row=4, max_row=akhir), Reference(ws, min_col=3, min_row=4, max_row=akhir), title="Strategi")
    seri.marker.symbol = "circle"
    seri.graphicalProperties.line.noFill = True
    sebar.series.append(seri)
    sebar.legend = None
    sebar.height, sebar.width = 9, 16
    ws.add_chart(sebar, "H3")

    # ------------------------------------------------------------------ Kelayakan
    ws = wb.create_sheet("Kelayakan")
    judul(ws, "Analisis kelayakan empat aspek", "Skor 1 (lemah) sampai 5 (kuat) adalah penilaian penulis atas bukti di kolom D; bobot dapat diganti.")
    kepala(ws, 3, ["Aspek", "Kriteria", "Lensa", "Bukti", "Skor", "Bobot", "Skor x bobot", "Syarat agar tetap layak"])
    tgl = p["terpilih"]
    kriteria = [
        ("Operasional", "Perubahan produk terbatas", "feasible", "Hanya dasar nilai pertanggungan jaminan gempa di lima provinsi yang berubah; jaminan kebakaran tetap nilai penuh.", 4, 0.08, "Berlaku saat polis baru dan perpanjangan"),
        ("Operasional", "Data lokasi polis tersedia", "feasible", "Zona tarif gempa sudah ditetapkan per kabupaten/kota, jadi alamat risiko sudah dicatat.", 4, 0.06, "Koordinat wajib untuk nilai pertanggungan di atas Rp50 miliar"),
        ("Operasional", "Penerimaan nasabah dan pialang", "desirable", "Premi gempa turun 40% menurut skala loss limit; kreditur dapat menuntut nilai penuh.", 3, 0.10, "Lapisan di atas batas ditawarkan lewat koasuransi dengan penanggung lain"),
        ("Operasional", "Sumber daya manusia", "feasible", "Tim tiga orang (aktuaris, analis bahaya, analis data).", 4, 0.04, "Pelatihan underwriter dan pemasaran"),
        ("Hukum dan regulasi", "Dasar nilai pertanggungan first loss", "feasible", "SEOJK 6/2017 Bagian VI angka 1 huruf a membolehkan first loss/sub limit basis untuk asuransi gempa.", 5, 0.10, "Nilai deklarasi sama dengan nilai sebenarnya"),
        ("Hukum dan regulasi", "Tarif premi", "feasible", "Skala premi polis loss limit ada di SEOJK 6/2017 Lampiran III Tabel III.B.", 5, 0.06, "Tarif zona tetap dipakai"),
        ("Hukum dan regulasi", "Target solvabilitas internal", "feasible", "POJK 26/2025 Pasal 3 meminta target paling rendah 120% dengan memperhitungkan profil risiko dan hasil stress test.", 5, 0.08, "Target dan dasarnya disetujui direksi dan dewan komisaris"),
        ("Hukum dan regulasi", "Batas investasi", "feasible", "Profil Konservatif memenuhi seluruh batas AYD Pasal 15 POJK 26/2025.", 5, 0.06, "Deposito paling banyak 20% per bank"),
        ("Teknologi", "Mesin simulasi", "feasible", "Model berbasis kejadian dan simulasi neraca sudah berjalan: 100.000 tahun dalam hitungan menit di komputer biasa, dengan R dan Python.", 4, 0.08, "Kode dan data diarsipkan; uji otomatis dijalankan tiap pembaruan"),
        ("Teknologi", "Data", "feasible", "Katalog USGS, statistik MAIPARK, dan statistik OJK terbuka dan diperbarui berkala.", 4, 0.05, "Data polis sendiri menggantikan sebaran industri"),
        ("Teknologi", "Ketepatan model", "feasible", "Uji balik menunjukkan model di bawah klaim nyata untuk gempa dekat sumber; kurva kerentanan memberi rentang lebar.", 2, 0.07, "Kalibrasi tahunan dan validasi independen; target memuat kurva terberat"),
        ("Modal", "Tambahan ekuitas", "viable", f"Tambahan {idn(tgl['ekuitas'] - s0['ekuitas'])} per 100 premi, {idn(100 * (tgl['ekuitas'] - s0['ekuitas']) / tgl['laba_bersih_rata'], 0)}% laba bersih setahun.", 5, 0.10, "Laba ditahan sampai penyangga terbentuk"),
        ("Modal", "Hemat modal", "viable", f"Ekuitas {idn(tgl['ekuitas'])} lawan {idn(p['modal_saja']['ekuitas'])} pada jalan modal saja.", 5, 0.07, "Batas ganti rugi tetap berlaku"),
        ("Modal", "Nilai ekonomi", "viable", "Nilai kini bersih terhadap jalan modal saja bergantung pada biaya modal; lihat lembar Sensitivitas.", 3, 0.05, "Biaya modal di atas titik impas"),
    ]
    for i, (aspek, kr, lensa, bukti, skor, bobot, syarat) in enumerate(kriteria, 4):
        tulis(ws, f"A{i}", aspek); tulis(ws, f"B{i}", kr, bungkus=True); tulis(ws, f"C{i}", lensa)
        tulis(ws, f"D{i}", bukti, bungkus=True); tulis(ws, f"E{i}", skor, F_MASUK, "0"); tulis(ws, f"F{i}", bobot, F_MASUK, PERSEN1)
        tulis(ws, f"G{i}", f"=E{i}*F{i}", F_BIASA, "0.00"); tulis(ws, f"H{i}", syarat, bungkus=True)
    ak = 3 + len(kriteria)
    tulis(ws, f"A{ak + 2}", "Ringkasan per aspek", F_TEBAL)
    kepala(ws, ak + 3, ["Aspek", "Jumlah bobot", "Skor rata-rata tertimbang", "Putusan"])
    for i, aspek in enumerate(["Operasional", "Hukum dan regulasi", "Teknologi", "Modal"], ak + 4):
        tulis(ws, f"A{i}", aspek)
        tulis(ws, f"B{i}", f'=SUMIF($A$4:$A${ak},A{i},$F$4:$F${ak})', F_BIASA, PERSEN1)
        tulis(ws, f"C{i}", f'=SUMIF($A$4:$A${ak},A{i},$G$4:$G${ak})/B{i}', F_BIASA, "0.00", ISI_HASIL)
        tulis(ws, f"D{i}", f'=IF(C{i}>=4,"Layak",IF(C{i}>=3,"Layak bersyarat","Belum layak"))')
    tulis(ws, f"A{ak + 8}", "Seluruh aspek", F_TEBAL)
    tulis(ws, f"B{ak + 8}", f"=SUM(F4:F{ak})", F_TEBAL, PERSEN1)
    tulis(ws, f"C{ak + 8}", f"=SUM(G4:G{ak})/B{ak + 8}", F_TEBAL, "0.00", ISI_HASIL)
    tulis(ws, f"D{ak + 8}", f'=IF(C{ak + 8}>=4,"Layak",IF(C{ak + 8}>=3,"Layak bersyarat","Belum layak"))', F_TEBAL)
    lebar(ws, [20, 30, 11, 70, 7, 9, 12, 52])

    # ------------------------------------------------------------------ Risiko
    ws = wb.create_sheet("Risiko")
    judul(ws, "Risiko pelaksanaan: Risk Matrix", "Kemungkinan dan dampak berskala 1 sampai 5 (penilaian penulis). Nilai = kemungkinan x dampak.")
    kepala(ws, 3, ["No", "Risiko", "Penyebab", "Kemungkinan", "Dampak", "Nilai", "Tingkat", "Tanggapan", "Pemilik"])
    risiko = [
        ("Nasabah besar menolak batas ganti rugi", "Kreditur mensyaratkan pertanggungan nilai penuh", 3, 4, "Lapisan di atas batas ditempatkan pada penanggung lain lewat koasuransi; mulai dari perpanjangan", "Underwriting"),
        ("Model meleset dari klaim nyata", "Kurva kerentanan dan eksposur per premi tidak pasti", 3, 5, "Target memuat kurva terberat; kalibrasi tahunan; validasi independen; rentang target 160% sampai 195%", "Aktuaria"),
        ("Premi hilang lebih besar dari perkiraan", "Nasabah memindahkan seluruh polis, bukan hanya jaminan gempa", 3, 3, "Pantau retensi per pialang; paket harga jaminan kebakaran", "Pemasaran"),
        ("Data lokasi tidak akurat", "Alamat tanpa koordinat; satu polis banyak lokasi", 3, 3, "Koordinat wajib untuk risiko besar; audit data per kuartal", "Operasi"),
        ("Gempa besar sebelum penyangga terbentuk", "Penyangga dibentuk dari laba tahun pertama", 1, 5, "Batas ganti rugi berlaku lebih dulu; pembagian laba ditunda", "Direksi"),
        ("Perubahan aturan permodalan", "Rancangan aturan modal baru berlaku bertahap mulai 2027", 2, 3, "Mesin simulasi memisahkan parameter regulasi; hitung ulang saat aturan terbit", "Kepatuhan"),
        ("Guncangan pasar bersamaan dengan bencana", "Penjualan aset saat harga turun", 2, 3, "Kantong likuiditas dari kas dan deposito; SBN bertenor pendek sampai menengah", "Investasi"),
    ]
    for i, (rs, sebab, km, dp_, tg, pm) in enumerate(risiko, 4):
        tulis(ws, f"A{i}", i - 3); tulis(ws, f"B{i}", rs, bungkus=True); tulis(ws, f"C{i}", sebab, bungkus=True)
        tulis(ws, f"D{i}", km, F_MASUK, "0"); tulis(ws, f"E{i}", dp_, F_MASUK, "0")
        tulis(ws, f"F{i}", f"=D{i}*E{i}", F_BIASA, "0")
        tulis(ws, f"G{i}", f'=IF(F{i}>=15,"Tinggi",IF(F{i}>=8,"Sedang","Rendah"))')
        tulis(ws, f"H{i}", tg, bungkus=True); tulis(ws, f"I{i}", pm)
    ar = 3 + len(risiko)
    tulis(ws, f"A{ar + 2}", "Jumlah risiko per tingkat", F_TEBAL)
    for i, tingkat in enumerate(["Tinggi", "Sedang", "Rendah"], ar + 3):
        tulis(ws, f"B{i}", tingkat); tulis(ws, f"D{i}", f'=COUNTIF($G$4:$G${ar},B{i})', F_BIASA, "0")
    lebar(ws, [5, 36, 40, 13, 9, 8, 10, 70, 14])

    # ------------------------------------------------------------------ Peta_Jalan_KPI
    ws = wb.create_sheet("Peta_Jalan_KPI")
    judul(ws, "Pelaksanaan: Roadmap dan KPI")
    kepala(ws, 3, ["Tahap", "Waktu", "Kegiatan", "Keluaran"])
    jalan = [
        ("1. Dasar", "Triwulan IV 2026", "Persetujuan direksi dan dewan komisaris; basis data eksposur; mesin simulasi dipasang", "Kebijakan Geo-Solvency Shield; data polis bergeokode"),
        ("2. Produk", "Triwulan I 2027", "Klausul first loss untuk jaminan gempa di zona inti; pedoman underwriting; pelatihan", "Klausul dan tarif siap dipakai"),
        ("3. Berlaku bertahap", "Triwulan II 2027", "Batas ganti rugi berlaku untuk polis baru dan perpanjangan risiko di atas Rp50 miliar di zona inti", "Sekitar 82% eksposur zona inti berbatas pada akhir siklus perpanjangan"),
        ("4. Berlaku penuh", "Triwulan III sampai IV 2027", "Seluruh risiko zona inti; kantong likuiditas dibentuk; laba ditahan sampai target tercapai", "Rasio solvabilitas pada target; kas dan deposito menutup klaim satu kejadian"),
        ("5. Kalibrasi", "Tiap tahun mulai 2028", "Model dijalankan ulang dengan data polis dan klaim terbaru; validasi independen; target ditinjau", "Laporan akumulasi dan uji skenario untuk direksi"),
    ]
    for i, b in enumerate(jalan, 4):
        for j, v in enumerate(b):
            tulis(ws, f"{'ABCD'[j]}{i}", v, bungkus=True)
    r = 4 + len(jalan) + 1
    tulis(ws, f"A{r}", "Indikator kinerja", F_TEBAL)
    kepala(ws, r + 1, ["Indikator", "Sekarang", "Sasaran", "Sumber ukur"])
    lik = h["terpilih"]["kantong_likuiditas"]
    kpi = [
        ("P(R < 100%) setahun, kurva campuran", f"={HM}B{R['peluang_di_bawah_100']}", "paling besar 0,50%", "Simulasi tahunan", PERSEN),
        ("P(R < 100%) setahun, kurva terberat", f"={HM}B{R['peluang_aulady']}", "paling besar 0,50%", "Simulasi tahunan", PERSEN),
        ("Rasio solvabilitas", f"={HM}B{R['target']}", f"paling rendah {idn(100 * s_pilih, 0)}%", "Laporan bulanan", PERSEN1),
        ("Rasio setelah skenario megathrust", f"={HM}B{R['rasio_setelah_megathrust']}", "paling rendah 100%", "Uji skenario per kuartal", PERSEN1),
        ("Eksposur zona inti yang berbatas", 0, "100%", "Basis data polis", PERSEN1),
        ("Kas dan deposito terhadap klaim satu kejadian 1 dalam 200 tahun", lik["K75"]["kas_dan_deposito"] / p["sekarang"]["klaim_kejadian_var_99_5"], "paling rendah 100%", "Laporan investasi", PERSEN1),
        ("Polis gempa dengan koordinat lokasi", "belum diukur", "paling rendah 95%", "Basis data polis", None),
    ]
    for i, (nama, kini, sasar, sumber, fmt) in enumerate(kpi, r + 2):
        tulis(ws, f"A{i}", nama, bungkus=True)
        tulis(ws, f"B{i}", kini, F_TAUT if isinstance(kini, str) and kini.startswith("=") else F_MASUK, fmt)
        tulis(ws, f"C{i}", sasar); tulis(ws, f"D{i}", sumber)
    lebar(ws, [44, 24, 76, 58])

    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        for baris_ in ws.iter_rows():
            for c in baris_:
                if c.value is not None and c.font is None:
                    c.font = F_BIASA
    wb.save(BERKAS)
    print("tersimpan:", BERKAS)


def _strategi(st, kolom):
    import pandas as pd
    df = pd.read_csv(DIR_T / "2_model_bencana" / "keluaran" / "04_strategi.csv").set_index("strategi")
    return df.loc[st, kolom]


if __name__ == "__main__":
    main()
